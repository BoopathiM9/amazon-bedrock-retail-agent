import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

def generate_docx():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    PRIMARY_COLOR = RGBColor(15, 23, 42)      # Slate 900
    ACCENT_COLOR = RGBColor(79, 70, 229)      # Indigo 600
    TEXT_DARK = RGBColor(30, 41, 59)          # Slate 800
    TEXT_MUTED = RGBColor(71, 85, 105)        # Slate 600

    def add_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.font.name = 'Arial'
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = ACCENT_COLOR
        
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="1" w:color="4F46E5"/></w:pBdr>')
        pPr.append(pBdr)

    # Header
    hp = doc.add_paragraph()
    hp.paragraph_format.space_after = Pt(1)
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = hp.add_run('BOOPATHI M\n')
    name_run.font.name = 'Arial'
    name_run.font.size = Pt(18)
    name_run.font.bold = True
    name_run.font.color.rgb = PRIMARY_COLOR

    sub_run = hp.add_run('Data & AI Intern @ Cloud Kinetics  |  B.Tech Artificial Intelligence & Data Science (2026 Batch)\n')
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(10)
    sub_run.font.bold = True
    sub_run.font.color.rgb = ACCENT_COLOR

    contact_run = hp.add_run('Chennai / Salem, Tamil Nadu, India  |  +91 8825991927  |  boopathi11std@gmail.com\nLinkedIn: linkedin.com/in/Boopathi-M-60341b269  |  Portfolio: boopathi-portfolio.vercel.app')
    contact_run.font.name = 'Arial'
    contact_run.font.size = Pt(9)
    contact_run.font.color.rgb = TEXT_MUTED

    # Professional Summary
    add_heading('Professional Summary')
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)
    s_run = sp.add_run(
        'Proactive AI & Data Science engineer with industry experience as a Data & AI Intern at Cloud Kinetics (Chennai). '
        'Specialized in Generative AI, Large Language Models (LLMs), AWS Bedrock, serverless architectures (AWS Lambda, S3), and Python. '
        'Proven expertise building autonomous AI agents, computer vision diagnostic pipelines, and scalable web applications. '
        'Certified with 3 NPTEL Elite credentials from IITs in Distributed Systems and Cloud Computing, eager to deliver high-impact AI products.'
    )
    s_run.font.name = 'Arial'
    s_run.font.size = Pt(9.5)
    s_run.font.color.rgb = TEXT_DARK

    # Work Experience
    add_heading('Work Experience & Internships')

    # Exp 1: Cloud Kinetics
    exp1 = doc.add_paragraph()
    exp1.paragraph_format.space_before = Pt(2)
    exp1.paragraph_format.space_after = Pt(1)
    r1 = exp1.add_run('Cloud Kinetics')
    r1.font.bold = True
    r1.font.size = Pt(10)
    r1_sub = exp1.add_run('  —  Data & AI Intern (Chennai Office)')
    r1_sub.font.bold = True
    r1_sub.font.size = Pt(9.5)
    r1_sub.font.color.rgb = ACCENT_COLOR
    r1_date = exp1.add_run('\tAugust 2026 – October 2026')
    r1_date.font.size = Pt(8.5)
    r1_date.font.italic = True
    r1_date.font.color.rgb = TEXT_MUTED

    pts1 = [
        'Selected for competitive Data & AI Internship at Cloud Kinetics Chennai office to research and develop enterprise AI applications.',
        'Engineered Generative AI workflows and agent architectures utilizing AWS Bedrock foundation models, prompt engineering, and API gateways.',
        'Developed serverless pipelines using AWS Lambda, S3, and Python for automated document ingestion, token management, and inference.',
        'Conducted rigorous cloud technical exercises, benchmarking model latency, hallucination rates, and multi-tenant data isolation.'
    ]
    for pt in pts1:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(0.5)
        r = bp.add_run(pt)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = TEXT_DARK

    # Exp 2: Optimus Technocrates
    exp2 = doc.add_paragraph()
    exp2.paragraph_format.space_before = Pt(3)
    exp2.paragraph_format.space_after = Pt(1)
    r2 = exp2.add_run('Optimus Technocrates')
    r2.font.bold = True
    r2.font.size = Pt(10)
    r2_sub = exp2.add_run('  —  Python & Data Analytics Intern')
    r2_sub.font.bold = True
    r2_sub.font.size = Pt(9.5)
    r2_sub.font.color.rgb = ACCENT_COLOR
    r2_date = exp2.add_run('\tJuly 2024 – August 2024')
    r2_date.font.size = Pt(8.5)
    r2_date.font.italic = True
    r2_date.font.color.rgb = TEXT_MUTED

    pts2 = [
        'Underwent comprehensive training in Python logic building, data cleansing, and exploratory data analysis (EDA).',
        'Built reusable Python modules to automate tabular data processing, summary analytics, and report formatting.'
    ]
    for pt in pts2:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(0.5)
        r = bp.add_run(pt)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.color.rgb = TEXT_DARK

    # Projects
    add_heading('Technical Projects')

    projects_data = [
        (
            'Amazon Bedrock Retail AI Agent',
            'AWS Bedrock, AWS Lambda, Python 3.11, REST APIs, Next.js',
            [
                'Constructed an autonomous retail agent utilizing AWS Bedrock foundation models and action groups for conversational commerce.',
                'Engineered serverless Lambda actions for automated catalog lookups, cart modifications, and real-time inventory queries.',
                'Implemented guardrails, structured JSON outputs, and fallback intents to eliminate hallucinations in transactional dialogue.'
            ]
        ),
        (
            'Lumbar Spine Abnormality Detection',
            'Python, OpenCV, TensorFlow/Keras, Image Preprocessing, NumPy',
            [
                'Engineered a computer vision classification pipeline detecting spinal deformities and abnormalities from radiographic X-ray scans.',
                'Applied adaptive histogram equalization, morphological filters, and feature extraction to assist early clinical diagnosis.'
            ]
        ),
        (
            'Synwell Care — Smart Wellness Monitoring Platform',
            'Python, Machine Learning, React, MySQL, REST API',
            [
                'Designed a predictive health tracking application analyzing vital signs and lifestyle parameters in real time.',
                'Constructed an interactive React dashboard with automated health alerts and lifestyle guidance based on ML predictive scoring.'
            ]
        )
    ]

    for p_title, p_tech, p_pts in projects_data:
        pp = doc.add_paragraph()
        pp.paragraph_format.space_before = Pt(2.5)
        pp.paragraph_format.space_after = Pt(0.5)
        t_run = pp.add_run(p_title)
        t_run.font.bold = True
        t_run.font.size = Pt(9.5)
        tech_run = pp.add_run(f'  |  {p_tech}')
        tech_run.font.size = Pt(8.5)
        tech_run.font.italic = True
        tech_run.font.color.rgb = TEXT_MUTED
        for pt in p_pts:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(0.5)
            r = bp.add_run(pt)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.font.color.rgb = TEXT_DARK

    # Technical Skills
    add_heading('Technical Skills')
    skills_data = [
        ('Generative AI & ML:', 'AWS Bedrock, LLMs, Prompt Engineering, Agentic AI, OpenCV, TensorFlow, Scikit-learn, Data Analytics'),
        ('Cloud & Backend:', 'AWS (Bedrock, Lambda, S3, IAM), Cloud Architecture, REST APIs, Git, GitHub'),
        ('Programming Languages:', 'Python, Java, SQL (MySQL), TypeScript, JavaScript, HTML5/CSS3'),
        ('Frameworks & Databases:', 'Next.js 15, React 19, Spring Boot, Node.js, MySQL, MongoDB'),
        ('Developer Tools:', 'VS Code, Eclipse IDE, Jupyter Notebook, Postman')
    ]
    for cat, val in skills_data:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0.5)
        c_run = sp.add_run(cat + ' ')
        c_run.font.bold = True
        c_run.font.size = Pt(9)
        c_run.font.color.rgb = TEXT_DARK
        v_run = sp.add_run(val)
        v_run.font.size = Pt(9)
        v_run.font.color.rgb = TEXT_MUTED

    # Education
    add_heading('Education')
    edus = [
        ('B.Tech in Artificial Intelligence & Data Science', 'Muthayammal Engineering College (Autonomous), Anna University', '2022 – 2026', 'CGPA: 7.8 / 10'),
        ('Higher Secondary Certificate (HSC - 12th)', 'SRV Matric Higher Secondary School, Attur (State Board)', '2020 – 2022', 'Score: 76%'),
        ('Secondary School Leaving Certificate (SSLC - 10th)', 'SRV Matric Higher Secondary School, Attur (State Board)', '2019 – 2020', 'Score: 86%')
    ]
    for deg, inst, yr, score in edus:
        ep = doc.add_paragraph()
        ep.paragraph_format.space_after = Pt(0.5)
        d_run = ep.add_run(deg + '  ')
        d_run.font.bold = True
        d_run.font.size = Pt(9)
        i_run = ep.add_run(inst + '  ')
        i_run.font.size = Pt(8.5)
        i_run.font.color.rgb = TEXT_MUTED
        s_run = ep.add_run(f'[{score}]  ({yr})')
        s_run.font.bold = True
        s_run.font.size = Pt(8.5)
        s_run.font.color.rgb = ACCENT_COLOR

    # Certifications
    add_heading('Certifications & Achievements')
    certs = [
        ('TATA Group & Forage:', 'Gen AI Powered Data Analytics Job Simulation'),
        ('NPTEL (IIT) Elite + Silver:', 'Foundation of Cloud IoT Edge ML'),
        ('NPTEL (IIT) Elite:', 'Distributed Systems'),
        ('NPTEL (IIT) Elite:', 'Fundamental Algorithms: Design and Analysis'),
        ('Conference Paper Presentation:', '"Artificial Intelligence & Machine Learning in Healthcare", Selvam Engineering College')
    ]
    for c_title, c_desc in certs:
        cp = doc.add_paragraph(style='List Bullet')
        cp.paragraph_format.space_after = Pt(0.5)
        t = cp.add_run(c_title + ' ')
        t.font.bold = True
        t.font.size = Pt(9)
        d = cp.add_run(c_desc)
        d.font.size = Pt(9)
        d.font.color.rgb = TEXT_DARK

    out_docx_public = r'C:\portfolio\boopathi-portfolio\public\Boopathi_Resume.docx'
    out_docx_downloads = r'C:\Users\boopa\Downloads\Boopathi_M_Resume_Updated.docx'
    doc.save(out_docx_public)
    doc.save(out_docx_downloads)
    print('Generated DOCX resumes successfully!')

def generate_pdf():
    pdf_path_public = r'C:\portfolio\boopathi-portfolio\public\Boopathi_Resume.pdf'
    pdf_path_downloads = r'C:\Users\boopa\Downloads\Boopathi_Resume_Updated.pdf'

    doc = SimpleDocTemplate(
        pdf_path_public,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=32,
        bottomMargin=32
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=20,
        alignment=1, # Center
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=2
    )

    sub_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#4F46E5'),
        spaceAfter=3
    )

    contact_style = ParagraphStyle(
        'DocContact',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#475569'),
        spaceAfter=8
    )

    section_heading = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        textColor=colors.HexColor('#4F46E5'),
        spaceBefore=6,
        spaceAfter=2
    )

    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=3
    )

    bullet_style = ParagraphStyle(
        'BulletStyle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#1E293B'),
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=1.5
    )

    role_style = ParagraphStyle(
        'RoleTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=11,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=1
    )

    elements = []

    # Header
    elements.append(Paragraph("BOOPATHI M", title_style))
    elements.append(Paragraph("Data & AI Intern @ Cloud Kinetics  |  B.Tech Artificial Intelligence & Data Science (2026)", sub_style))
    elements.append(Paragraph(
        "Chennai / Salem, Tamil Nadu, India &nbsp;|&nbsp; +91 8825991927 &nbsp;|&nbsp; boopathi11std@gmail.com<br/>"
        "<b>LinkedIn:</b> linkedin.com/in/Boopathi-M-60341b269 &nbsp;|&nbsp; <b>Portfolio:</b> boopathi-portfolio.vercel.app",
        contact_style
    ))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=4, spaceBefore=0))

    # Summary
    elements.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
    elements.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#4F46E5'), spaceAfter=4, spaceBefore=0))
    elements.append(Paragraph(
        "Proactive AI & Data Science engineer with real-world experience as a <b>Data & AI Intern at Cloud Kinetics (Chennai)</b>. "
        "Specialized in Generative AI, Large Language Models (LLMs), AWS Bedrock, serverless architectures (AWS Lambda, S3), and Python. "
        "Demonstrated track record designing autonomous multi-step agents, computer vision diagnostic pipelines, and responsive web applications. "
        "Recipient of 3 NPTEL Elite certifications from IITs, committed to architecting scalable, enterprise-grade AI systems.",
        body_style
    ))

    # Experience
    elements.append(Paragraph("WORK EXPERIENCE & INTERNSHIPS", section_heading))
    elements.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#4F46E5'), spaceAfter=4, spaceBefore=0))

    # Cloud Kinetics
    exp1_table_data = [[
        Paragraph("<b>Cloud Kinetics</b> — <font color='#4F46E5'><b>Data & AI Intern (Chennai Office)</b></font>", role_style),
        Paragraph("<font color='#64748B'><i>August 2026 – October 2026</i></font>", ParagraphStyle('RAlign', parent=body_style, alignment=2))
    ]]
    t1 = Table(exp1_table_data, colWidths=[380, 160])
    t1.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOTTOMPADDING', (0,0), (-1,-1), 1), ('TOPPADDING', (0,0), (-1,-1), 0)]))
    elements.append(t1)

    elements.append(Paragraph("• Selected for competitive Data & AI Internship at Cloud Kinetics Chennai office to engineer enterprise AI and cloud solutions.", bullet_style))
    elements.append(Paragraph("• Developed Generative AI workflows and autonomous agents utilizing AWS Bedrock foundation models, prompt pipelines, and APIs.", bullet_style))
    elements.append(Paragraph("• Built serverless compute pipelines with AWS Lambda, S3, and Python for automated document parsing, token tracking, and inference.", bullet_style))
    elements.append(Paragraph("• Executed hands-on cloud technical exercises, benchmarking latency, hallucination rates, and multi-tenant cloud architecture.", bullet_style))
    elements.append(Spacer(1, 3))

    # Optimus Technocrates
    exp2_table_data = [[
        Paragraph("<b>Optimus Technocrates</b> — <font color='#4F46E5'><b>Python & Data Analytics Intern</b></font>", role_style),
        Paragraph("<font color='#64748B'><i>July 2024 – August 2024</i></font>", ParagraphStyle('RAlign2', parent=body_style, alignment=2))
    ]]
    t2 = Table(exp2_table_data, colWidths=[380, 160])
    t2.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('BOTTOMPADDING', (0,0), (-1,-1), 1), ('TOPPADDING', (0,0), (-1,-1), 0)]))
    elements.append(t2)

    elements.append(Paragraph("• Completed immersive hands-on training in Python algorithmic logic building, data cleansing, and exploratory data analysis (EDA).", bullet_style))
    elements.append(Paragraph("• Programmed modular Python software packages to automate data cleaning, statistical evaluation, and real-time report generation.", bullet_style))

    # Projects
    elements.append(Paragraph("TECHNICAL PROJECTS", section_heading))
    elements.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#4F46E5'), spaceAfter=4, spaceBefore=0))

    elements.append(Paragraph("<b>Amazon Bedrock Retail AI Agent</b> &nbsp;|&nbsp; <i>AWS Bedrock, AWS Lambda, Python 3.11, REST APIs, Next.js</i>", role_style))
    elements.append(Paragraph("• Built an autonomous retail agent utilizing AWS Bedrock foundation models and action groups for conversational commerce.", bullet_style))
    elements.append(Paragraph("• Implemented Lambda Action Groups to query product catalogs, verify real-time stock levels, and dynamically process orders.", bullet_style))
    elements.append(Paragraph("• Configured prompt safety guardrails and structured JSON parsing to guarantee reliable, hallucination-free retail recommendations.", bullet_style))
    elements.append(Spacer(1, 2))

    elements.append(Paragraph("<b>Lumbar Spine Abnormality Detection</b> &nbsp;|&nbsp; <i>Python, OpenCV, TensorFlow/Keras, Image Preprocessing</i>", role_style))
    elements.append(Paragraph("• Developed a computer vision AI framework to detect spinal deformities and abnormalities from medical radiographic X-ray scans.", bullet_style))
    elements.append(Paragraph("• Applied adaptive histogram equalization, morphological filtering, and contour segmentation to assist early diagnostic screening.", bullet_style))
    elements.append(Spacer(1, 2))

    elements.append(Paragraph("<b>Synwell Care — Smart Wellness Monitoring Platform</b> &nbsp;|&nbsp; <i>Python, Machine Learning, React, MySQL</i>", role_style))
    elements.append(Paragraph("• Designed an intelligent health vitals monitoring web platform analyzing personal biomarkers and lifestyle data in real time.", bullet_style))
    elements.append(Paragraph("• Engineered an interactive React dashboard with automated health risk alerts and lifestyle recommendations using ML predictive models.", bullet_style))

    # Skills
    elements.append(Paragraph("TECHNICAL SKILLS", section_heading))
    elements.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#4F46E5'), spaceAfter=4, spaceBefore=0))
    elements.append(Paragraph("<b>Generative AI & ML:</b> AWS Bedrock, Large Language Models (LLMs), Prompt Engineering, Agentic Workflows, OpenCV, TensorFlow, Scikit-learn", body_style))
    elements.append(Paragraph("<b>Cloud & Backend:</b> AWS (Bedrock, Lambda, S3, IAM), Cloud Architecture, REST APIs, Git, GitHub", body_style))
    elements.append(Paragraph("<b>Programming Languages:</b> Python, Java, SQL (MySQL), TypeScript, JavaScript, HTML5/CSS3", body_style))
    elements.append(Paragraph("<b>Frameworks & Databases:</b> Next.js 15, React 19, Spring Boot, Node.js, MySQL, MongoDB", body_style))
    elements.append(Paragraph("<b>Developer Tools:</b> VS Code, Eclipse IDE, Jupyter Notebook, Postman", body_style))

    # Education
    elements.append(Paragraph("EDUCATION", section_heading))
    elements.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#4F46E5'), spaceAfter=4, spaceBefore=0))
    elements.append(Paragraph("<b>B.Tech in Artificial Intelligence & Data Science</b> &nbsp;|&nbsp; Muthayammal Engineering College (Autonomous), Anna Univ &nbsp; <b>[CGPA: 7.8]</b> &nbsp; (2022–2026)", body_style))
    elements.append(Paragraph("<b>Higher Secondary Certificate (HSC)</b> &nbsp;|&nbsp; SRV Matric Higher Secondary School, Attur (State Board) &nbsp; <b>[Score: 76%]</b> &nbsp; (2020–2022)", body_style))
    elements.append(Paragraph("<b>Secondary School Leaving Certificate (SSLC)</b> &nbsp;|&nbsp; SRV Matric Higher Secondary School, Attur (State Board) &nbsp; <b>[Score: 86%]</b> &nbsp; (2019–2020)", body_style))

    # Certifications
    elements.append(Paragraph("CERTIFICATIONS & ACHIEVEMENTS", section_heading))
    elements.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#4F46E5'), spaceAfter=4, spaceBefore=0))
    elements.append(Paragraph("• <b>TATA Group & Forage:</b> Gen AI Powered Data Analytics Job Simulation", bullet_style))
    elements.append(Paragraph("• <b>NPTEL (IIT) Elite + Silver:</b> Foundation of Cloud IoT Edge ML", bullet_style))
    elements.append(Paragraph("• <b>NPTEL (IIT) Elite:</b> Distributed Systems", bullet_style))
    elements.append(Paragraph("• <b>NPTEL (IIT) Elite:</b> Fundamental Algorithms: Design and Analysis", bullet_style))
    elements.append(Paragraph("• <b>Conference Paper Presentation:</b> \"Artificial Intelligence & Machine Learning in Healthcare\", Selvam Engineering College", bullet_style))

    doc.build(elements)

    # Also copy to downloads
    import shutil
    shutil.copyfile(pdf_path_public, pdf_path_downloads)
    print('Generated PDF resumes successfully!')

if __name__ == '__main__':
    generate_docx()
    generate_pdf()
